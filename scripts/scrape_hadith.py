"""
Generic scraper for hadith collections on sunnah.com.

Usage:
    python scrape_hadith.py <collection>

Examples:
    python scrape_hadith.py malik
    python scrape_hadith.py bukhari
"""

import os
import sys
import time
import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

BASE_URL = "https://sunnah.com"
HEADERS = {"User-Agent": "Mozilla/5.0"}
REQUEST_DELAY = 1  # seconds between requests


def fetch_page(url):
    """Fetch a page with retry logic."""
    for attempt in range(4):
        try:
            resp = requests.get(url, headers=HEADERS, timeout=30)
            resp.raise_for_status()
            return resp.text
        except requests.RequestException as e:
            if attempt < 3:
                wait = 2 ** (attempt + 1)
                print(f"  Retry {attempt + 1}/3 after {wait}s due to: {e}")
                time.sleep(wait)
            else:
                raise


def scrape_books(collection):
    """Scrape the list of all books from the main collection page."""
    url = f"{BASE_URL}/{collection}"
    print(f"Fetching book list from {url}")
    html = fetch_page(url)
    soup = BeautifulSoup(html, "lxml")

    books = []
    for entry in soup.find_all("div", class_="book_title"):
        num_el = entry.find("div", class_="book_number")
        eng_el = entry.find("div", class_="english_book_name")
        arb_el = entry.find("div", class_="arabic_book_name")

        book_num = num_el.get_text(strip=True) if num_el else ""
        eng_title = eng_el.get_text(strip=True) if eng_el else ""
        arb_title = arb_el.get_text(strip=True) if arb_el else ""

        link = entry.find("a")
        href = link.get("href", "") if link else f"/{collection}/{book_num}"

        books.append({
            "number": book_num,
            "title_en": eng_title,
            "title_ar": arb_title,
            "url": f"{BASE_URL}{href}",
        })

    print(f"Found {len(books)} books")
    return books


def scrape_hadith(book_url):
    """Scrape all hadith from a single book page."""
    html = fetch_page(book_url)
    soup = BeautifulSoup(html, "lxml")

    hadith_list = []
    containers = soup.find_all("div", class_="actualHadithContainer")

    for container in containers:
        # Hadith number from anchor name attribute
        anchor = container.find("a", attrs={"name": True})
        hadith_num = anchor.get("name", "") if anchor else ""

        # English text
        eng_div = container.find("div", class_="english_hadith_full")
        eng_text = eng_div.get_text(separator="\n", strip=True) if eng_div else ""

        # Arabic text
        arb_div = container.find("div", class_="arabic_hadith_full")
        arb_text = arb_div.get_text(separator="\n", strip=True) if arb_div else ""

        # Reference from the table
        ref_text = ""
        table = container.find("table")
        if table:
            cells = table.find_all("td")
            if len(cells) >= 2:
                ref_text = cells[1].get_text(strip=True).lstrip(": ").replace("\xa0", " ").strip()

        hadith_list.append({
            "number": hadith_num,
            "reference": ref_text,
            "english": eng_text,
            "arabic": arb_text,
        })

    return hadith_list


def create_excel(books_data, output_path):
    """Create an Excel workbook with the scraped data."""
    wb = Workbook()

    # ── Styles ──
    header_font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="2E6B4E", end_color="2E6B4E", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    arabic_font = Font(name="Arial", size=11)
    arabic_alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
    english_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    center_alignment = Alignment(horizontal="center", vertical="top")

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # ── Books Index Sheet ──
    ws_index = wb.active
    ws_index.title = "Books Index"

    index_headers = ["Book Number", "English Title", "Arabic Title", "Hadith Count"]
    for col, header in enumerate(index_headers, 1):
        cell = ws_index.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for row, (book, hadith_list) in enumerate(books_data, 2):
        ws_index.cell(row=row, column=1, value=int(book["number"])).alignment = center_alignment
        ws_index.cell(row=row, column=2, value=book["title_en"]).alignment = english_alignment
        cell_ar = ws_index.cell(row=row, column=3, value=book["title_ar"])
        cell_ar.font = arabic_font
        cell_ar.alignment = arabic_alignment
        ws_index.cell(row=row, column=4, value=len(hadith_list)).alignment = center_alignment
        for col in range(1, 5):
            ws_index.cell(row=row, column=col).border = thin_border

    ws_index.column_dimensions["A"].width = 14
    ws_index.column_dimensions["B"].width = 35
    ws_index.column_dimensions["C"].width = 35
    ws_index.column_dimensions["D"].width = 14

    # ── All Hadith Sheet ──
    ws_all = wb.create_sheet("All Hadith")
    all_headers = [
        "Book Number", "Book Title (EN)", "Book Title (AR)",
        "Hadith Number", "Reference", "English Hadith", "Arabic Hadith",
    ]
    for col, header in enumerate(all_headers, 1):
        cell = ws_all.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    current_row = 2
    for book, hadith_list in books_data:
        for h in hadith_list:
            ws_all.cell(row=current_row, column=1, value=int(book["number"])).alignment = center_alignment
            ws_all.cell(row=current_row, column=2, value=book["title_en"]).alignment = english_alignment
            cell_bar = ws_all.cell(row=current_row, column=3, value=book["title_ar"])
            cell_bar.font = arabic_font
            cell_bar.alignment = arabic_alignment
            ws_all.cell(row=current_row, column=4, value=h["number"]).alignment = center_alignment
            ws_all.cell(row=current_row, column=5, value=h["reference"]).alignment = center_alignment
            ws_all.cell(row=current_row, column=6, value=h["english"]).alignment = english_alignment
            cell_har = ws_all.cell(row=current_row, column=7, value=h["arabic"])
            cell_har.font = arabic_font
            cell_har.alignment = arabic_alignment

            for col in range(1, 8):
                ws_all.cell(row=current_row, column=col).border = thin_border

            current_row += 1

    ws_all.column_dimensions["A"].width = 12
    ws_all.column_dimensions["B"].width = 30
    ws_all.column_dimensions["C"].width = 30
    ws_all.column_dimensions["D"].width = 14
    ws_all.column_dimensions["E"].width = 20
    ws_all.column_dimensions["F"].width = 80
    ws_all.column_dimensions["G"].width = 80

    # ── Freeze top row on both sheets ──
    ws_index.freeze_panes = "A2"
    ws_all.freeze_panes = "A2"

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    print(f"\nExcel file saved to: {output_path}")


def main():
    if len(sys.argv) < 2:
        print("Usage: python scrape_hadith.py <collection>")
        print("Example: python scrape_hadith.py malik")
        sys.exit(1)

    collection = sys.argv[1].lower()
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_dir = os.path.dirname(script_dir)
    output_path = os.path.join(project_dir, "output", collection, f"{collection}_hadith.xlsx")

    books = scrape_books(collection)
    books_data = []
    total_hadith = 0

    for i, book in enumerate(books):
        print(f"[{i+1}/{len(books)}] Scraping Book {book['number']}: {book['title_en']}...", end=" ")
        hadith_list = scrape_hadith(book["url"])
        print(f"{len(hadith_list)} hadith")
        books_data.append((book, hadith_list))
        total_hadith += len(hadith_list)

        if i < len(books) - 1:
            time.sleep(REQUEST_DELAY)

    print(f"\nTotal: {len(books)} books, {total_hadith} hadith")
    create_excel(books_data, output_path)


if __name__ == "__main__":
    main()
